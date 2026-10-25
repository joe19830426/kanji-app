"""小学校 漢字小テスト自動作成アプリ（縦書き・A4/B4ヨコ・印刷対応）

起動:
    pip install streamlit anthropic python-docx
    streamlit run app.py

設定 (.streamlit/secrets.toml または環境変数):
    APP_PASSWORD      = "好きなパスワード"   # 未設定なら認証なしで動作
    ANTHROPIC_API_KEY = "sk-ant-..."         # 問題文の自動作成・画像/PDFの読み取りに使用
"""

import base64
import hmac
import html
import io
import json
import math
import os
import random
import re

import streamlit as st
import streamlit.components.v1 as components

MODEL = "claude-sonnet-5-5"

PUBLISHERS = ["光村図書", "東京書籍", "教育出版", "学校図書", "三省堂", "その他"]
TERMS = ["1学期", "2学期", "3学期"]
COUNTS = [10, 20, 30, 40, 50]
PAPERS = {
    "A4ヨコ": (297, 210, "297mm 210mm"),
    "B4ヨコ": (364, 257, "364mm 257mm"),
}
KINDS = ["標準", "お助け（2画ヒント）", "ルビ付き", "お助け＋ルビ"]

# ---------------------------------------------------------------------------
# 学年別漢字配当表（小学校学習指導要領 平成29年告示・全1,026字）
# 1年80字 / 2年160字 / 3年200字 / 4年202字 / 5年193字 / 6年191字
# ---------------------------------------------------------------------------
GRADE_KANJI = {
    1: "一右雨円王音下火花貝学気九休玉金空月犬見五口校左三山子四糸字耳七車手十出女小上森人水正生青夕石赤千川先早草足村大男竹中虫町天田土二日入年白八百文木本名目立力林六",
    2: "引羽雲園遠何科夏家歌画回会海絵外角楽活間丸岩顔汽記帰弓牛魚京強教近兄形計元言原戸古午後語工公広交光考行高黄合谷国黒今才細作算止市矢姉思紙寺自時室社弱首秋週春書少場色食心新親図数西声星晴切雪船線前組走多太体台地池知茶昼長鳥朝直通弟店点電刀冬当東答頭同道読内南肉馬売買麦半番父風分聞米歩母方北毎妹万明鳴毛門夜野友用曜来里理話",
    3: "悪安暗医委意育員院飲運泳駅央横屋温化荷界開階寒感漢館岸起期客究急級宮球去橋業曲局銀区苦具君係軽血決研県庫湖向幸港号根祭皿仕死使始指歯詩次事持式実写者主守取酒受州拾終習集住重宿所暑助昭消商章勝乗植申身神真深進世整昔全相送想息速族他打対待代第題炭短談着注柱丁帳調追定庭笛鉄転都度投豆島湯登等動童農波配倍箱畑発反坂板皮悲美鼻筆氷表秒病品負部服福物平返勉放味命面問役薬由油有遊予羊洋葉陽様落流旅両緑礼列練路和",
    4: "愛案以衣位茨印英栄媛塩岡億加果貨課芽賀改械害街各覚潟完官管関観願岐希季旗器機議求泣給挙漁共協鏡競極熊訓軍郡群径景芸欠結建健験固功好香候康佐差菜最埼材崎昨札刷察参産散残氏司試児治滋辞鹿失借種周祝順初松笑唱焼照城縄臣信井成省清静席積折節説浅戦選然争倉巣束側続卒孫帯隊達単置仲沖兆低底的典伝徒努灯働特徳栃奈梨熱念敗梅博阪飯飛必票標不夫付府阜富副兵別辺変便包法望牧末満未民無約勇要養浴利陸良料量輪類令冷例連老労録",
    5: "圧囲移因永営衛易益液演応往桜可仮価河過快解格確額刊幹慣眼紀基寄規喜技義逆久旧救居許境均禁句型経潔件険検限現減故個護効厚耕航鉱構興講告混査再災妻採際在財罪殺雑酸賛士支史志枝師資飼示似識質舎謝授修述術準序招証象賞条状常情織職制性政勢精製税責績接設絶祖素総造像増則測属率損貸態団断築貯張停提程適統堂銅導得毒独任燃能破犯判版比肥非費備評貧布婦武復複仏粉編弁保墓報豊防貿暴脈務夢迷綿輸余容略留領歴",
    6: "胃異遺域宇映延沿恩我灰拡革閣割株干巻看簡危机揮貴疑吸供胸郷勤筋系敬警劇激穴券絹権憲源厳己呼誤后孝皇紅降鋼刻穀骨困砂座済裁策冊蚕至私姿視詞誌磁射捨尺若樹収宗就衆従縦縮熟純処署諸除承将傷障蒸針仁垂推寸盛聖誠舌宣専泉洗染銭善奏窓創装層操蔵臓存尊退宅担探誕段暖値宙忠著庁頂腸潮賃痛敵展討党糖届難乳認納脳派拝背肺俳班晩否批秘俵腹奮並陛閉片補暮宝訪亡忘棒枚幕密盟模訳郵優預幼欲翌乱卵覧裏律臨朗論",
}

# ---------------------------------------------------------------------------
# 画数表（学年別漢字配当表の「画数から探す」表に基づく）と、学期ごとの目安
#   学期ごとの新出漢字は教科書会社・学校によって異なり、公的な割り当てはありません。
#   ここでは学年の漢字を「画数の少ない順」に3等分した“目安”として扱います。
# ---------------------------------------------------------------------------
STROKE_TABLE = {
    1: {1: "一", 2: "九七十二人入八力", 3: "下口三山子女小上夕千川大土", 4: "円王火月犬五手水中天日文木六",
        5: "右玉左四出正生石田白本目立", 6: "気休糸字耳先早竹虫年百名", 7: "花貝見車赤足村男町",
        8: "雨学金空青林", 9: "音草", 10: "校", 12: "森"},
    2: {2: "刀", 3: "丸弓工才万", 4: "引牛元戸午公今止少心切太内父分方毛友", 5: "外兄古広市矢台冬半母北用",
        6: "羽回会交光考行合寺自色西多地池当同肉米毎", 7: "何角汽近形言谷作社図声走体弟売麦来里",
        8: "画岩京国姉知長直店東歩妹明門夜", 9: "科海活計後思室首秋春食星前茶昼点南風",
        10: "夏家記帰原高紙時弱書通馬", 11: "魚強教黄黒細週雪船組鳥野理", 12: "雲絵間場晴朝答道買番",
        13: "園遠楽新数電話", 14: "歌語算読聞鳴", 15: "線", 16: "親頭", 18: "顔曜"},
    3: {2: "丁", 4: "化区反予", 5: "央去号皿仕写主申世他打代皮氷平由礼",
        6: "安曲血向死次式守州全有羊両列", 7: "医究局君決住助身対投豆坂返役",
        8: "委育泳岸苦具幸使始事実者取受所昔注定波板表服物放味命油和",
        9: "屋界客急級係研県指持拾重昭乗神相送待炭柱追度畑発美秒品負面洋",
        10: "員院荷起宮庫根酒消真息速庭島配倍病勉流旅", 11: "悪球祭終習宿商章深進族第帳笛転都動部問",
        12: "飲運温開階寒期軽湖港歯集暑勝植短着湯登等童悲筆遊葉陽落", 13: "暗意感漢業詩想鉄農福路",
        14: "駅銀鼻様緑練", 15: "横談調箱", 16: "館橋整薬", 18: "題"},
    4: {4: "欠氏井不夫", 5: "以加功札司失必付辺包末未民令", 6: "衣印各共好成争仲兆伝灯老",
        7: "位改完岐希求芸佐材児初臣折束沖低努阪兵別利良冷労",
        8: "英岡果芽官季泣協径固刷参治周松卒底的典奈念府阜法牧例",
        9: "茨栄軍建香昨祝城信省浅単栃飛変便約勇要", 10: "案害挙訓郡候差残借笑席倉孫帯徒特梅浴料連",
        11: "貨械健康菜埼崎産鹿唱清巣側梨敗票副望陸", 12: "媛賀街覚給極景結最散滋順焼然隊達博飯富満無量",
        13: "愛塩群試辞照節戦続置働", 14: "管関旗漁熊察種静説徳", 15: "億課潟器縄選熱標養輪",
        16: "機積録", 18: "観験類", 19: "願鏡", 20: "議競"},
    5: {3: "久士", 4: "支比仏", 5: "圧永可刊旧句史示犯布弁", 6: "因仮件再在団任",
        7: "囲応快技均告災志似序条状判防余", 8: "易往価河居効妻枝舎述招制性毒版肥非武",
        9: "紀逆型限故厚査政祖則独保迷", 10: "益桜格個耕航財殺師修素造能破粉脈容留",
        11: "移液眼基寄規救許経険現混採授術常情責接設率断張停堂得貧婦務略",
        12: "営過喜検減証象税絶測属貸貯提程統費備評復報貿", 13: "解幹義禁鉱罪資飼準勢損墓豊夢",
        14: "演慣境構際雑酸精製総像増態適銅複綿領歴", 15: "確潔賛質賞導編暴", 16: "衛興築燃輸",
        17: "講謝績", 18: "額織職", 19: "識", 20: "護"},
    6: {3: "干己寸亡", 4: "尺収仁片", 5: "穴冊処庁幼", 6: "宇灰危机吸后至舌存宅", 7: "我系孝困私否批忘乱卵",
        8: "延沿拡供券呼刻若宗承垂担宙忠届乳拝並宝枚", 9: "胃映革巻看皇紅砂姿宣専泉洗染奏退段派背肺律",
        10: "恩株胸降骨座蚕射従純除将針値展討党納俳班秘俵陛朗", 11: "異域郷済視捨推盛窓探著頂脳閉訪密訳郵欲翌",
        12: "割揮貴勤筋敬裁策詞就衆善創装尊痛晩補棒", 13: "絹源署傷蒸聖誠暖腸賃腹幕盟預裏",
        14: "閣疑誤穀誌磁障銭層認暮模", 15: "遺劇権熟諸蔵誕潮敵論", 16: "激憲鋼樹縦操糖奮",
        17: "厳縮優覧", 18: "簡難臨", 19: "警臓"},
}

KANJI_STROKES = {c: n for tbl in STROKE_TABLE.values() for n, cs in tbl.items() for c in cs}

RANGE_TERM = "この学期に習う漢字（目安）"
RANGE_UNTIL = "この学期までに習う漢字（目安）"
RANGE_ALL = "学年すべて"


def term_kanji(grade: int, term_idx: int):
    """学年の漢字を画数の少ない順に3等分し、学期ごとの目安にする（配当表の順で返す）。"""
    chars = list(GRADE_KANJI[grade])
    order = sorted(range(len(chars)), key=lambda i: (KANJI_STROKES.get(chars[i], 99), i))
    rank = {chars[i]: r for r, i in enumerate(order)}
    k = len(chars)
    b1, b2 = round(k / 3), round(k * 2 / 3)
    lo, hi = [(0, b1), (b1, b2), (b2, k)][term_idx]
    return [c for c in chars if lo <= rank[c] < hi]


def kanji_range(grade: int, term: str, how: str):
    ti = TERMS.index(term)
    if how == RANGE_ALL:
        return list(GRADE_KANJI[grade])
    if how == RANGE_UNTIL:
        s = set()
        for t in range(ti + 1):
            s |= set(term_kanji(grade, t))
        return [c for c in GRADE_KANJI[grade] if c in s]
    return term_kanji(grade, ti)


# ---------------------------------------------------------------------------
# サンプル問題（APIキー未設定時の代替・補充用） (答えの語, よみ, 例文)
#   例文の {W} が答えの位置。[漢|かん] は「ルビ付き」のときだけルビになる。
# ---------------------------------------------------------------------------
SAMPLE = {
    1: [
        ("学校", "がっこう", "ぼくの{W}は、[山|やま]の[下|した]だ。"),
        ("花火", "はなび", "みんなで{W}を[見|み]た。"),
        ("森林", "しんりん", "{W}の[中|なか]は、しずかだ。"),
        ("青空", "あおぞら", "{W}に[白|しろ]い雲がうかぶ。"),
        ("先生", "せんせい", "{W}に[本|ほん]をかりた。"),
        ("足音", "あしおと", "うしろから{W}がした。"),
        ("草花", "くさばな", "{W}に[水|みず]をやる。"),
        ("町", "まち", "[大|おお]きな{W}をあるく。"),
        ("天気", "てんき", "あしたも{W}がいい。"),
        ("左右", "さゆう", "{W}を[見|み]てわたる。"),
        ("見える", "みえる", "とおくの[山|やま]が{W}。"),
        ("休む", "やすむ", "きょうは[学校|がっこう]を{W}。"),
    ],
    2: [
        ("親切", "しんせつ", "{W}にして、よろこばれた。"),
        ("元気", "げんき", "あさから{W}いっぱいだ。"),
        ("海外", "かいがい", "{W}の[国|くに]に[行|い]く。"),
        ("通学", "つうがく", "{W}で[電車|でんしゃ]にのる。"),
        ("図書", "としょ", "{W}[室|しつ]で[本|ほん]を[読|よ]む。"),
        ("計算", "けいさん", "{W}のテストをする。"),
        ("家", "いえ", "[大|おお]きな{W}にすむ。"),
        ("黒", "くろ", "{W}い[犬|いぬ]がいる。"),
        ("東西", "とうざい", "{W}に[長|なが]い[道|みち]がある。"),
        ("会話", "かいわ", "[友|とも]だちと{W}をする。"),
        ("答える", "こたえる", "[先生|せんせい]に{W}。"),
        ("遠い", "とおい", "[学校|がっこう]は{W}。"),
    ],
    3: [
        ("運動", "うんどう", "[毎朝|まいあさ]{W}をする。"),
        ("意見", "いけん", "[自分|じぶん]の{W}を[言|い]う。"),
        ("係", "かかり", "{W}の[仕事|しごと]をする。"),
        ("決心", "けっしん", "{W}して、[勉強|べんきょう]する。"),
        ("湖", "みずうみ", "{W}の[水|みず]がすんでいる。"),
        ("研究", "けんきゅう", "{W}を[進|すす]める。"),
        ("階", "かい", "[二|に]{W}の[教室|きょうしつ]へ[行|い]く。"),
        ("整理", "せいり", "[本|ほん]だなを{W}する。"),
        ("写真", "しゃしん", "[家族|かぞく]の{W}をとる。"),
        ("勉強", "べんきょう", "[算数|さんすう]を{W}する。"),
        ("決める", "きめる", "[係|かかり]を{W}。"),
        ("住む", "すむ", "[町|まち]に{W}。"),
    ],
    4: [
        ("博物館", "はくぶつかん", "[町|まち]の{W}に[行|い]く。"),
        ("最近", "さいきん", "{W}、[本|ほん]をよく[読|よ]む。"),
        ("共同", "きょうどう", "みんなで{W}して[作|つく]る。"),
        ("努力", "どりょく", "{W}をつづける。"),
        ("信号", "しんごう", "{W}が[赤|あか]にかわる。"),
        ("給食", "きゅうしょく", "{W}のじかんだ。"),
        ("観察", "かんさつ", "[花|はな]を{W}する。"),
        ("改札", "かいさつ", "[駅|えき]の{W}を[通|とお]る。"),
        ("希望", "きぼう", "[大|おお]きな{W}をもつ。"),
        ("季節", "きせつ", "[春|はる]は、いい{W}だ。"),
        ("浅い", "あさい", "[川|かわ]が{W}。"),
        ("働く", "はたらく", "[父|ちち]は[会社|かいしゃ]で{W}。"),
    ],
    5: [
        ("経験", "けいけん", "[自分|じぶん]の{W}を[話|はな]す。"),
        ("資料", "しりょう", "{W}を[集|あつ]める。"),
        ("規則", "きそく", "{W}を[守|まも]る。"),
        ("可能", "かのう", "{W}かどうか[考|かんが]える。"),
        ("在庫", "ざいこ", "[店|みせ]の{W}を[調|しら]べる。"),
        ("現在", "げんざい", "{W}の[時間|じかん]を[知|し]る。"),
        ("情報", "じょうほう", "{W}を[調|しら]べる。"),
        ("移動", "いどう", "[遠|とお]くへ{W}する。"),
        ("構成", "こうせい", "{W}を[考|かんが]える。"),
        ("提案", "ていあん", "[新|あたら]しい{W}をする。"),
        ("任せる", "まかせる", "[仕事|しごと]を{W}。"),
        ("確かめる", "たしかめる", "[答|こた]えを{W}。"),
    ],
    6: [
        ("宇宙", "うちゅう", "{W}を[旅|たび]する。"),
        ("貴重", "きちょう", "{W}な[資料|しりょう]を[見|み]る。"),
        ("郵便", "ゆうびん", "{W}で[手紙|てがみ]を[送|おく]る。"),
        ("認識", "にんしき", "[正|ただ]しく{W}する。"),
        ("誠実", "せいじつ", "{W}な[人|ひと]だ。"),
        ("縦横", "じゅうおう", "{W}に[線|せん]を[引|ひ]く。"),
        ("宝物", "たからもの", "{W}を[見|み]つけた。"),
        ("展覧会", "てんらんかい", "{W}に[行|い]く。"),
        ("憲法", "けんぽう", "{W}を[学|まな]ぶ。"),
        ("磁石", "じしゃく", "{W}で[鉄|てつ]をくっつける。"),
        ("訪ねる", "たずねる", "[友|とも]だちを{W}。"),
        ("忘れる", "わすれる", "[宿題|しゅくだい]を{W}。"),
    ],
}

KANJI_RE = re.compile(r"[\u4e00-\u9fff々]")


def sample_pool(grade):
    """その学年までの内蔵サンプル問題"""
    return [{"word": w, "reading": r, "sentence": s}
            for g in range(1, grade + 1) for w, r, s in SAMPLE[g]]


# ---------------------------------------------------------------------------
# 認証（APP_PASSWORD を設定した場合のみ）
# ---------------------------------------------------------------------------
def secret(name: str) -> str:
    """secrets.toml → 環境変数 の順で設定値を読む（secrets.toml が無くても落ちない）"""
    try:
        v = st.secrets.get(name, "")
    except Exception:
        v = ""
    return str(v or os.environ.get(name, "")).strip()


def check_password() -> bool:
    if st.session_state.get("authenticated"):
        return True
    expected = secret("APP_PASSWORD")
    if not expected:
        return True
    st.title("🔒 漢字小テスト作成")
    with st.form("login"):
        pw = st.text_input("パスワード", type="password")
        ok = st.form_submit_button("ログイン")
    if ok:
        if hmac.compare_digest(pw.encode(), str(expected).encode()):
            st.session_state["authenticated"] = True
            st.rerun()
        else:
            st.error("パスワードが違います。")
    return False


# ---------------------------------------------------------------------------
# Claude API
# ---------------------------------------------------------------------------
def api_key() -> str:
    return secret("ANTHROPIC_API_KEY")


def call_claude(content, max_tokens=8000) -> str:
    import anthropic
    key = api_key()
    if not key:
        raise RuntimeError("ANTHROPIC_API_KEY が設定されていません。")
    client = anthropic.Anthropic(api_key=key, timeout=150.0, max_retries=1)   # 無限に待たない
    resp = client.messages.create(model=MODEL, max_tokens=max_tokens,
                                  messages=[{"role": "user", "content": content}])
    return "".join(b.text for b in resp.content if b.type == "text")


def parse_json_list(text: str):
    """AIの応答から {...} を1件ずつ読み取る。前置き・コードブロック・途中切れ・壊れた1件があっても、
    読めたものだけを返す（文中の [漢|かん] の ] に惑わされない）。"""
    dec = json.JSONDecoder()
    items, pos = [], 0
    while True:
        i = text.find("{", pos)
        if i < 0:
            break
        try:
            obj, end = dec.raw_decode(text, i)
        except ValueError:
            pos = i + 1
            continue
        if isinstance(obj, dict):
            items.append(obj)
        pos = end
    if not items:
        raise ValueError("AIの応答から問題を読み取れませんでした。もう一度お試しください。")
    return items


def generate_with_claude(grade, kanji, n, exclude=None):
    prompt = f"""小学{grade}年生向けの漢字書き取りテストの問題を、ちょうど{n}問作ってください。
出題に使う漢字（この中から選ぶ）: {kanji}

【ルール】
1. 各問は「答えとなる語（漢字で書く）」を1つ決める。答えの語には、上の出題漢字を少なくとも1字含める。
2. 答えの語に使う漢字・文中の漢字は、小学{grade}年生までに習う漢字だけにする。習っていない漢字はひらがなにする。
3. 熟語（2字以上）、1字の語、送り仮名のある語（例：任せる・細かい）をバランスよく混ぜる。送り仮名のある語は2割ほど。
4. 例文は{grade}年生が読める短い文（答えを除いて15字以内）。答えの部分を {{W}} と書く。
5. 文中の答え以外の漢字は [漢字|よみ] の形でひらがなの読みを付ける（例：[先生|せんせい]に{{W}}。）。
6. 出題漢字がなるべく偏らないようにし、同じ語を重複させない。易しい問題から順に並べる。
JSON配列だけを出力（説明やコードブロックは不要）:
[{{"kanji":"出題漢字1字","word":"答えの語","reading":"ひらがな","sentence":"例文"}}]"""
    if exclude:
        prompt += "\n\n次の語はすでに作成済みなので使わないこと: " + "、".join(exclude)
    rows = parse_json_list(call_claude(prompt))
    allowed = set("".join(GRADE_KANJI[g] for g in range(1, grade + 1)))
    out = []
    for r in rows:
        w, s = str(r.get("word", "")).strip(), str(r.get("sentence", ""))
        if not w or "{W}" not in s:
            continue
        if any(KANJI_RE.match(c) and c not in allowed for c in w):
            continue
        out.append({"word": w, "reading": str(r.get("reading", "")).strip(), "sentence": s})
    return out


# ---------------------------------------------------------------------------
# 添付プリントの読み取り
# ---------------------------------------------------------------------------
MARK_RE = re.compile(r"【([^|｜】]+)(?:[|｜]([^】]+))?】")
NUM_RE = re.compile(r"^\s*[\(（]?[0-9０-９①-⑳]+[\)）.．、]?\s*")


def parse_text(text: str):
    """行ごとに【語|よみ】を含む文、または「語,よみ,例文」形式を読み取る。"""
    entries = []
    for raw in text.splitlines():
        line = NUM_RE.sub("", raw.strip())
        if not line:
            continue
        m = MARK_RE.search(line)
        if m:
            entries.append({"word": m.group(1).strip(), "reading": (m.group(2) or "").strip(),
                            "sentence": line[:m.start()] + "{W}" + line[m.end():]})
            continue
        parts = [p.strip() for p in re.split(r"[,，\t]", line)]
        if len(parts) >= 3 and parts[0] and parts[0] in parts[2]:
            entries.append({"word": parts[0], "reading": parts[1],
                            "sentence": parts[2].replace(parts[0], "{W}", 1)})
    return entries


def decode_bytes(b: bytes) -> str:
    for enc in ("utf-8-sig", "cp932", "utf-16"):
        try:
            return b.decode(enc)
        except Exception:
            continue
    return b.decode("utf-8", errors="ignore")


def read_with_claude(data: bytes, media_type: str):
    b64 = base64.standard_b64encode(data).decode()
    kind = "document" if media_type == "application/pdf" else "image"
    block = {"type": kind, "source": {"type": "base64", "media_type": media_type, "data": b64}}
    prompt = ("添付は小学校の漢字テストのプリントです。書き問題（ひらがなを漢字に直す問題）を全て読み取り、"
              "読み問題の場合は下線部の語を答えの語として同様に扱ってください。"
              "次のJSON配列だけを出力してください（説明やコードブロックは不要）。\n"
              '[{"sentence":"文全体。答えの漢字語の部分だけを {W} に置き換える","word":"答えの漢字語","reading":"ひらがなのよみ"}]')
    rows = parse_json_list(call_claude([block, {"type": "text", "text": prompt}]))
    return [{"word": r["word"], "reading": r.get("reading", ""), "sentence": r["sentence"]}
            for r in rows if r.get("word") and "{W}" in r.get("sentence", "")]


def read_attachment(f):
    name, data = f.name.lower(), f.getvalue()
    if name.endswith((".txt", ".csv", ".md")):
        return parse_text(decode_bytes(data))
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        lines = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                lines.append(",".join(c.text for c in row.cells))
        return parse_text("\n".join(lines))
    if name.endswith(".pdf"):
        return read_with_claude(data, "application/pdf")
    if name.endswith((".png", ".jpg", ".jpeg")):
        return read_with_claude(data, "image/png" if name.endswith(".png") else "image/jpeg")
    raise ValueError("対応していない形式です。")


# ---------------------------------------------------------------------------
# 添付プリントの文の作り変え・ルビ付け（答えの語・よみは絶対に変えない）
# ---------------------------------------------------------------------------
THEMES = ["学校", "家", "公園", "町", "季節", "家族", "友だち", "食べ物", "動物", "天気",
          "スポーツ", "旅行", "係の仕事", "図書室", "のりもの", "夏休み", "運動会"]


def plain_text(sentence: str) -> str:
    """[漢|かん] 形式のルビ指定を外した文"""
    return RUBY_RE.sub(lambda m: m.group(1), sentence)


def rewrite_sentences(items, grade, seed):
    """答えの語（漢字・送り仮名・よみ）は元のまま、例文だけをAIで作り直す。(items, 失敗数) を返す。"""
    rng = random.Random(seed)
    themes = "・".join(rng.sample(THEMES, 3))
    allowed = set("".join(GRADE_KANJI[g] for g in range(1, grade + 1)))
    for e in items:
        allowed |= {c for c in e["word"] if KANJI_RE.match(c)}
    payload = [{"i": i, "word": e["word"], "reading": e.get("reading", ""),
                "old": e["sentence"].replace("{W}", e["word"])} for i, e in enumerate(items)]
    prompt = f"""小学{grade}年生向けの漢字テストの例文を作り直してください。
各問には「答えの語」があります。答えの語は絶対に変えず、例文だけを新しく作ります。

【ルール】
1. 新しい例文では、答えの語が入る位置を {{W}} と書き、{{W}} は1回だけ使う。答えの語そのもの（漢字・送り仮名・活用形）は変えない。
   例：答えが「任せる」なら {{W}} は「任せる」の形で使う（「任せた」「任せて」にしない）。
   例：答えが「土」なら、{{W}}の前後に字を足して「土地」「土曜」のような別の語にしない。
2. 答えの語を {{W}} 以外の場所に書かない。
3. 元の文（old）とは違う場面・言い回しにする。今回の場面のヒント：{themes}。問ごとに変化をつけ、同じ文型を続けない。
4. {grade}年生が読める自然な短い文（答えを除いて18字以内）。句点（。）で終える。
5. 文中の答え以外の漢字は、小学{grade}年生までに習う漢字だけを使い、[漢字|ひらがなの読み] の形で読みを付ける
   （例：[先生|せんせい]と{{W}}。）。習っていない漢字はひらがなで書く。
JSON配列だけを出力（説明やコードブロックは不要）: [{{"i":0,"sentence":"..."}}]

入力: {json.dumps(payload, ensure_ascii=False)}"""
    rows = parse_json_list(call_claude(prompt))
    out = [dict(e) for e in items]
    for r in rows:
        try:
            i, s = int(r["i"]), str(r["sentence"]).strip()
        except Exception:
            continue
        if not (0 <= i < len(items)) or s.count("{W}") != 1:
            continue
        word = items[i]["word"]
        text = plain_text(s.replace("{W}", ""))
        if word in text:                                   # 答えの語が別の場所にも出ている
            continue
        if any(KANJI_RE.match(c) and c not in allowed for c in text):   # 未習の漢字
            continue
        out[i]["sentence"] = s                             # word / reading は元のまま（AIの出力は使わない）
    changed = sum(1 for a, b in zip(items, out) if a["sentence"] != b["sentence"])
    return out, len(items) - changed


def annotate_ruby(items):
    """ルビ指定のない文に、AIで [漢字|よみ] を付ける。文字そのものが変わった場合は元の文を使う。"""
    todo = [(i, e) for i, e in enumerate(items)
            if "[" not in e["sentence"] and KANJI_RE.search(e["sentence"].replace("{W}", ""))]
    if not todo or not api_key():
        return items
    payload = [{"i": i, "sentence": e["sentence"]} for i, e in todo]
    prompt = f"""次の各文の {{W}} 以外の漢字に、ルビ用の読みを [漢字|ひらがなの読み] の形で付けてください。
熟語は1語まとめて付ける（例：[先生|せんせい]）。{{W}} の部分と、文の文字は一切変更・追加・削除しない。
JSON配列だけを出力（説明やコードブロックは不要）: [{{"i":0,"sentence":"..."}}]

入力: {json.dumps(payload, ensure_ascii=False)}"""
    out = [dict(e) for e in items]
    for r in parse_json_list(call_claude(prompt)):
        try:
            i, s = int(r["i"]), str(r["sentence"])
        except Exception:
            continue
        if 0 <= i < len(items) and plain_text(s) == items[i]["sentence"]:
            out[i]["sentence"] = s
    return out


# ---------------------------------------------------------------------------
# 問題の作成
# ---------------------------------------------------------------------------
def build_items(grade, pool, selected, n, seed, use_random):
    """学年の漢字から問題を作る。(items, notes) を返す。"""
    rng = random.Random(seed)
    notes = []
    if use_random:
        pool_k = rng.sample(pool, min(n, len(pool)))
    else:
        pool_k = list(selected)
        if len(pool_k) > n:
            pool_k = rng.sample(pool_k, n)
    kanji = "".join(pool_k)

    items = []
    if api_key():
        try:
            for _ in range(3):                      # 足りない分は、作成済みの語を除いて追加で依頼する
                need = n - len(items)
                if need <= 0:
                    break
                new = generate_with_claude(grade, kanji, need, exclude=[e["word"] for e in items])
                have = {e["word"] for e in items}
                added = [e for e in new if e["word"] not in have]
                if not added:
                    break
                items += added
            if len(items) < n:
                notes.append(f"AIが作成できた問題は{len(items)}問でした。残りは内蔵のサンプル問題で補います。")
        except Exception as ex:
            notes.append(f"AIによる問題作成の途中で失敗しました（{ex}）。足りない分はサンプル問題で補います。")
    else:
        notes.append("ANTHROPIC_API_KEY が未設定のため、内蔵のサンプル問題から作成しました。")

    if len(items) < n:
        ks = set(pool_k)
        used = {e["word"] for e in items}
        allp = [e for e in sample_pool(grade) if e["word"] not in used]
        matched = [e for e in allp if any(c in ks for c in e["word"])]
        others = [e for e in allp if e not in matched]
        rng.shuffle(matched)
        rng.shuffle(others)
        need = n - len(items)
        if len(matched) < need and not use_random:
            notes.append("選んだ漢字を含む内蔵サンプルが少ないため、他のサンプル問題で補いました。")
        items += (matched + others)[:need]
    if len(items) < n:
        notes.append(f"問題数が足りず、{len(items)}問のプリントになりました。")
    return items[:n], notes


def build_from_attachment(att, n, use_all, shuffle, rewrite, grade, seed, ruby):
    """添付プリントの問題から作る。並びはそのまま／ランダム、文は元のまま／作り変え。"""
    rng = random.Random(seed)
    notes = []
    items = [dict(e) for e in att]
    if shuffle:
        rng.shuffle(items)
    limit = min(len(items), 50) if use_all else min(n, len(items))
    if use_all and len(items) > 50:
        notes.append(f"添付プリントは{len(items)}問ありますが、プリントに収まる50問までを使います。")
    items = items[:limit]
    notes.append(f"添付プリントの問題を{len(items)}問使いました"
                 f"（{'ランダムな並び' if shuffle else 'プリント通りの並び'}）。")
    if not use_all and n > len(att):
        notes.append(f"添付プリントの問題数（{len(att)}問）が、選んだ問題数より少ないため、{len(att)}問になりました。")

    if rewrite:
        if not api_key():
            notes.append("ANTHROPIC_API_KEY が未設定のため、文は作り変えず元の文を使いました。")
        else:
            try:
                items, failed = rewrite_sentences(items, grade, seed)
                notes.append("答えの漢字はそのまま、文を作り直しました。"
                             + (f"（{failed}問は条件に合う文にできず、元の文を使っています）" if failed else ""))
            except Exception as ex:
                notes.append(f"文の作り直しに失敗したため、元の文を使いました（{ex}）")
    if ruby and api_key():
        try:
            items = annotate_ruby(items)
        except Exception as ex:
            notes.append(f"ルビの自動付けに失敗しました（{ex}）")
    elif ruby:
        notes.append("ルビを自動で付けるには ANTHROPIC_API_KEY が必要です（添付の文にルビ指定がある場合のみ表示されます）。")
    return items, notes


def load_attachment(up):
    """添付ファイルの読み取り結果をキャッシュして、再作成のたびにAIを呼ばないようにする。"""
    key = (up.name, up.size)
    cache = st.session_state.get("att_cache")
    if cache and cache["key"] == key:
        return cache["items"]
    items = read_attachment(up)
    st.session_state["att_cache"] = {"key": key, "items": items}
    return items


# ---------------------------------------------------------------------------
# プリントHTML生成（縦書き）
# ---------------------------------------------------------------------------
CSS = """
*{box-sizing:border-box}
html,body{margin:0}
body{background:#e9ecef;color:#111;
 font-family:"UD デジタル 教科書体 NK-R","UD Digi Kyokasho NK-R","Hiragino Mincho ProN","Yu Mincho","BIZ UDMincho","Noto Serif CJK JP","Noto Serif JP","IPAexMincho",serif}
.toolbar{position:sticky;top:0;z-index:5;display:flex;gap:8px;align-items:center;padding:8px;
 background:#fff;border-bottom:1px solid #ccc;font-family:sans-serif;font-size:13px;color:#555}
.toolbar button{font-size:16px;padding:10px 18px;min-height:44px;border:0;border-radius:8px;background:#1a73e8;color:#fff}
.scaler{transform-origin:top left;padding:8px}
.sheet{width:__W__mm;height:__H__mm;background:#fff;padding:5mm 6mm;margin:0 0 8px;display:flex;
 flex-direction:column;box-shadow:0 1px 6px #0003;overflow:hidden}
.body{flex:1;min-height:0;display:grid;direction:rtl;grid-template-columns:auto 1fr;gap:3mm}

/* 右端：テスト名／組・番／名前（1行）＋説明・得点 */
.names{direction:ltr;writing-mode:vertical-rl;height:100%}
.line{display:flex;inline-size:100%;margin-left:1.5mm}
.line1{block-size:11.5mm;flex-direction:row}
.line2{block-size:8mm;flex-direction:row}
.tt{flex:none;font-size:5mm;font-weight:bold;letter-spacing:.4mm;line-height:2.1;white-space:nowrap;padding-top:.5mm}
.tt.long{font-size:4mm}
.nf{display:flex;flex:none;border:.4mm solid #111;font-size:4.2mm;line-height:1.1}
.nf .lb{flex:none;text-align:center;padding:1.2mm 0;letter-spacing:.8mm}
.nf .sp{flex:1}
.nf.cls{block-size:11.5mm;inline-size:auto;height:36mm;margin-top:2mm}
.nf.name{block-size:11.5mm;flex:1;margin-top:2mm}
.desc{flex:1;font-size:3.2mm;line-height:1.7;color:#222}
.score{flex:none;border:.4mm solid #111;height:30mm;block-size:8mm;font-size:3.6mm;display:flex;margin-top:2mm}
.score span{padding:1mm 0;text-align:center;flex:none}
.score .sp{flex:1}

/* 問題エリア */
.qs{direction:rtl;display:grid;min-height:0;min-width:0;border-top:.25mm solid #888;border-right:.25mm solid #888}
.q{direction:ltr;writing-mode:vertical-rl;display:flex;flex-wrap:wrap;align-content:flex-start;align-items:center;
 padding:1.2mm 1mm;font-size:var(--fs);
 line-height:var(--lh);border-left:.25mm solid #888;border-bottom:.25mm solid #888;min-height:0;min-width:0}
.no{display:inline-block;width:5.4mm;height:5.4mm;border:.3mm solid #111;border-radius:50%;font-size:3.2mm;
 line-height:5mm;text-align:center;font-family:sans-serif;margin-inline-end:1.2mm;flex:none}
.u{flex:none}
.tcy{text-combine-upright:all;-webkit-text-combine:horizontal}
rt{font-size:.5em;line-height:1}

/* 解答欄：漢字ごとに□を連結／送り仮名つきは〔 〕枠。よみは枠の外（右）に配置 */
.answ{display:block;flex:none;position:relative;padding-right:calc(var(--fs)*.72);margin:.6mm 0}
.bxs{display:inline-flex}
.bx{display:block;width:var(--bx);height:var(--bx);border:.4mm solid #111;text-align:center;
 line-height:calc(var(--bx) - .8mm);font-size:calc(var(--bx)*.78);color:#c4001a;
 background:
  linear-gradient(90deg,transparent calc(50% - .1mm),#cfcfcf 0 calc(50% + .1mm),transparent 0),
  linear-gradient(transparent calc(50% - .1mm),#cfcfcf 0 calc(50% + .1mm),transparent 0)}
.bx+.bx{margin-top:-.4mm}
.brk{display:flex;align-items:center;justify-content:center;width:var(--bx);color:#c4001a;
 font-size:calc(var(--bx)*.78);letter-spacing:0;line-height:1;
 border-top:.45mm solid #111;border-bottom:.45mm solid #111;
 background:
  linear-gradient(#111,#111) 0 0/.45mm 2.6mm no-repeat,
  linear-gradient(#111,#111) 100% 0/.45mm 2.6mm no-repeat,
  linear-gradient(#111,#111) 0 100%/.45mm 2.6mm no-repeat,
  linear-gradient(#111,#111) 100% 100%/.45mm 2.6mm no-repeat}
.rd{position:absolute;right:0;top:0;bottom:0;display:flex;align-items:center;justify-content:center;
 font-size:calc(var(--fs)*.62);line-height:1;white-space:nowrap;color:#111}

/* お助けヒント */
.hint{display:block;flex:none;margin-top:.8mm;border:.3mm dashed #666;border-radius:1mm;padding:.5mm;vertical-align:top}
.hk{display:block;width:calc(var(--bx)*.6);height:calc(var(--bx)*.6);line-height:calc(var(--bx)*.6);
 font-size:calc(var(--bx)*.5);text-align:center;color:#444;
 -webkit-mask-image:linear-gradient(135deg,#000 0%,#000 40%,transparent 56%);
 mask-image:linear-gradient(135deg,#000 0%,#000 40%,transparent 56%)}

@page{size:__PAPER__ landscape;margin:0}
@media print{
 body{background:#fff}
 .toolbar{display:none}
 .scaler{transform:none!important;padding:0;width:auto!important}
 .sheet{height:calc(__H__mm - 1mm);margin:0;box-shadow:none;break-after:page;page-break-after:always}
 .sheet:last-child{break-after:auto;page-break-after:auto}
 *{-webkit-print-color-adjust:exact;print-color-adjust:exact}
}
"""

JS = """
function fitSheets(){
 var list=document.querySelectorAll('.qs');
 for(var q=0;q<list.length;q++){
  var qs=list[q],cells=qs.querySelectorAll('.q');
  if(!cells.length||qs.clientHeight<20)continue;
  var bx0=parseFloat(qs.getAttribute('data-bx')),fs0=parseFloat(qs.getAttribute('data-fs'));
  var hi=parseFloat(qs.getAttribute('data-max')),lo=0.5;
  var apply=function(s){
   qs.style.setProperty('--bx',(bx0*s).toFixed(2)+'mm');
   qs.style.setProperty('--fs',(fs0*(s<=1?Math.pow(s,0.6):1+(s-1)*0.2)).toFixed(2)+'mm');};
  var fits=function(){
   for(var i=0;i<cells.length;i++){var c=cells[i];
    if(c.scrollWidth>c.clientWidth+1||c.scrollHeight>c.clientHeight+1)return false;}
   return true;};
  apply(hi);
  if(!fits()){
   for(var k=0;k<14;k++){var mid=(lo+hi)/2;apply(mid);if(fits())lo=mid;else hi=mid;}
   apply(lo);
  }
  qs.setAttribute('data-fit',qs.style.getPropertyValue('--bx'));
 }
}
function fit(){var s=document.querySelector('.sheet'),c=document.querySelector('.scaler');
 if(!s||!s.offsetWidth)return;
 c.style.transform='none';var w=s.offsetWidth+16,k=Math.min(1,window.innerWidth/w);
 c.style.transform='scale('+k+')';}
function all(){fitSheets();fit();}
window.addEventListener('load',all);window.addEventListener('resize',all);
if(document.fonts&&document.fonts.ready){document.fonts.ready.then(all);}
window.onbeforeprint=function(){document.querySelector('.scaler').style.transform='none'};
window.onafterprint=all;
"""

RUBY_RE = re.compile(r"\[([^\]|]+)\|([^\]]+)\]")


def tate(s: str) -> str:
    """縦書き用：半角数字を縦中横にする"""
    return re.sub(r"\d{1,3}", lambda m: f'<span class="tcy">{m.group()}</span>',
                  html.escape(s, quote=False))


def has_okurigana(word: str) -> bool:
    return any(not KANJI_RE.match(c) for c in word)


def answer_html(e, answer: bool) -> str:
    word = e["word"]
    rd = f'<span class="rd">{html.escape(e["reading"], quote=False)}</span>' if e.get("reading") else ""
    if has_okurigana(word):
        inner = html.escape(word, quote=False) if answer else ""
        return (f'<span class="answ"><span class="brk" style="height:calc(var(--bx)*{len(word)})">'
                f'{inner}</span>{rd}</span>')
    boxes = "".join(f'<span class="bx">{html.escape(c, quote=False) if answer else ""}</span>' for c in word)
    return f'<span class="answ"><span class="bxs">{boxes}</span>{rd}</span>'


def text_units(text: str, ruby: bool) -> str:
    """文を1文字（ルビ付き語は1まとまり）ずつの部品にする。折り返しを正確に制御するため。"""
    def chars(t):
        # 句読点・閉じ括弧は直前の文字と一緒に折り返す（行頭に孤立させない）
        units = []
        for c in t:
            if units and c in "。、，．）」』！？":
                units[-1] += c
            else:
                units.append(c)
        return "".join(f'<span class="u">{html.escape(u, quote=False)}</span>' for u in units)
    out, pos = [], 0
    for m in RUBY_RE.finditer(text):
        out.append(chars(text[pos:m.start()]))
        if ruby:
            out.append(f'<span class="u"><ruby>{html.escape(m.group(1), quote=False)}'
                       f'<rt>{html.escape(m.group(2), quote=False)}</rt></ruby></span>')
        else:
            out.append(chars(m.group(1)))
        pos = m.end()
    out.append(chars(text[pos:]))
    return "".join(out)


def sentence_html(e, ruby: bool, answer: bool) -> str:
    pre, sep, post = e["sentence"].partition("{W}")
    return text_units(pre, ruby) + answer_html(e, answer) + text_units(post, ruby)


def cell_html(i, e, row, col, help_on, ruby, answer) -> str:
    no = f'<span class="no"><span class="tcy">{i}</span></span>'
    hint = ""
    if help_on and not answer:
        hint = f'<span class="hint"><span class="hk">{html.escape(e["word"][0], quote=False)}</span></span>'
    return (f'<div class="q" style="grid-row:{row};grid-column:{col}">'
            f'{no}{sentence_html(e, ruby, answer)}{hint}</div>')


def layout_for(n):
    """(各行の問題数リスト, 解答枠mm, 文字mm)
    枠の大きさは“ふつう”の基準値。実際の大きさは、JSが枠に収まる最大まで自動で調整する。"""
    if n <= 5:
        rows, bx, fs = 1, 20.0, 6.0
    elif n <= 10:
        rows, bx, fs = 2, 18.0, 5.4
    elif n <= 20:
        rows, bx, fs = 3, 14.5, 4.6
    elif n <= 30:
        rows, bx, fs = 3, 12.5, 4.2
    elif n <= 40:
        rows, bx, fs = 5, 10.5, 3.8
    else:
        rows, bx, fs = 4, 9.5, 3.5          # 50問は 13・13・12・12 の4行
    base, extra = divmod(n, rows)
    counts = [base + 1] * extra + [base] * (rows - extra)
    return counts, bx, fs


def sheet_html(items, title, kind, answer, publisher, term, max_scale=1.0) -> str:
    help_on = kind in ("お助け（2画ヒント）", "お助け＋ルビ")
    ruby = kind in ("ルビ付き", "お助け＋ルビ")
    n = len(items)
    counts, bx, fs = layout_for(n)
    bands, cols = len(counts), max(counts)
    lh = 2.05 if ruby else 1.6
    pts = f"{100 / n:.1f}".rstrip("0").rstrip(".")

    cells, idx = [], 0
    for r, c in enumerate(counts, start=1):
        for k in range(1, c + 1):
            cells.append(cell_html(idx + 1, items[idx], r, k, help_on, ruby, answer))
            idx += 1

    ttl = title + ("（解答）" if answer else "（100点満点）")
    long = " long" if len(ttl) > 20 else ""
    desc = ("赤字が答えです。" if answer else
            f"文中の　線のひらがなを漢字で書きましょう。（1問{pts}点）")
    meta = f"{publisher}・{term}"
    names = (
        '<div class="names">'
        f'<div class="line line1"><div class="tt{long}">{tate(ttl)}</div>'
        '<div class="nf cls"><span class="sp"></span><span class="lb">組</span><span class="sp"></span>'
        '<span class="lb">番</span><span class="sp"></span></div>'
        '<div class="nf name"><span class="lb">名前</span><span class="sp"></span></div></div>'
        f'<div class="line line2"><div class="desc">{tate(meta)}　{tate(desc)}</div>'
        '<div class="score"><span>得点</span><span class="sp"></span></div></div>'
        '</div>'
    )
    return (f'<section class="sheet"><div class="body">{names}'
            f'<div class="qs" data-bx="{bx}" data-fs="{fs}" data-max="{max_scale}" '
            f'style="--fs:{fs}mm;--bx:{bx}mm;--lh:{lh};'
            f'grid-template-columns:repeat({cols},1fr);grid-template-rows:repeat({bands},1fr)">'
            f'{"".join(cells)}</div></div></section>')


def document_html(sheets, paper) -> str:
    w, h, size = PAPERS[paper]
    css = CSS.replace("__W__", str(w)).replace("__H__", str(h)).replace("__PAPER__", size)
    return (f'<!doctype html><html lang="ja"><head><meta charset="utf-8">'
            f'<meta name="viewport" content="width=device-width,initial-scale=1"><style>{css}</style></head><body>'
            f'<div class="toolbar"><button onclick="window.print()">🖨 共有・プリント</button>'
            f'<span>{paper}・横向きで印刷（倍率100%／余白なし）</span></div>'
            f'<div class="scaler">{"".join(sheets)}</div><script>{JS}</script></body></html>')


# ---------------------------------------------------------------------------
# 画面
# ---------------------------------------------------------------------------
SIZES = {"自動（学年に合わせる）": None, "ふつう": 1.0, "大きめ": 1.2, "特大": 1.4}
AUTO_SCALE = {1: 1.4, 2: 1.3, 3: 1.2, 4: 1.1, 5: 1.0, 6: 1.0}
ORDERS = ["プリント通りの並び", "毎回ランダムに並び替える"]
COUNT_MODES = ["添付プリントの全問題を使う", "上で選んだ問題数に合わせる"]
PICK_MODES = ["一覧から選ぶ", "この範囲からランダムに選ぶ"]


def set_sel(key, chars):
    st.session_state[key] = list(chars)


def kanji_picker(key, pool):
    b1, b2, b3 = st.columns([1, 1, 3])
    b1.button("全選択", key=f"all_{key}", on_click=set_sel, args=(key, pool))
    b2.button("全解除", key=f"none_{key}", on_click=set_sel, args=(key, []))
    if hasattr(st, "pills"):
        sel = st.pills("使う漢字", pool, selection_mode="multi", key=key, label_visibility="collapsed")
    else:
        sel = st.multiselect("使う漢字", pool, key=key, label_visibility="collapsed")
    sel = sel or []
    b3.caption(f"{len(pool)}字のうち {len(sel)}字を選択中")
    return sel


def main():
    st.set_page_config(page_title="漢字小テスト作成", page_icon="📝", layout="wide")
    if not check_password():
        return
    st.markdown("<style>.stButton>button,.stDownloadButton>button{min-height:46px}"
                "div[role=radiogroup]{gap:.4rem}</style>", unsafe_allow_html=True)
    st.title("📝 漢字小テスト作成（縦書き）")

    st.session_state.setdefault("seed", random.randint(0, 10**6))

    with st.expander("⚙ 設定", expanded=True):
        c1, c2, c3, c4 = st.columns(4)
        grade = c1.selectbox("学年", [1, 2, 3, 4, 5, 6], format_func=lambda g: f"{g}年生")
        publisher = c2.selectbox("教科書会社", PUBLISHERS)
        term = c3.selectbox("学期", TERMS)
        paper = c4.selectbox("用紙", list(PAPERS))
        c5, c6, c7 = st.columns([1, 2, 2])
        n = c5.selectbox("問題数", COUNTS, index=1, format_func=lambda x: f"{x}問")
        size_opt = c6.selectbox("解答マスの大きさ", list(SIZES))
        title = c7.text_input("プリント名", value=f"{grade}年 漢字小テスト")
        kind = st.radio("プリントの種類", KINDS, horizontal=True)
        st.caption("解答マスは枠に収まる最大の大きさに自動調整されます。問題数が少ないほど大きくなるので、"
                   "1年生などで大きくしたいときは問題数を減らしてください。")

    with st.expander("🔤 使う漢字を選ぶ（学年別漢字配当表・学期ごとの目安つき）", expanded=True):
        c8, c9 = st.columns(2)
        pick = c8.radio("漢字の選び方", PICK_MODES, horizontal=True)
        rng_how = c9.radio("漢字の範囲", [RANGE_TERM, RANGE_UNTIL, RANGE_ALL])
        pool = kanji_range(grade, term, rng_how)
        st.caption(f"{grade}年生・{term}の範囲：{len(pool)}字（{grade}年生の漢字は全部で{len(GRADE_KANJI[grade])}字）。"
                   "※学期ごとの新出漢字は教科書会社・学校によって違い、公的な割り当てはありません。"
                   "ここでは画数の少ない順に3等分した目安で表示しています。実際の範囲と違うときは"
                   "「学年すべて」から選んでください。")
        selected = []
        if pick == PICK_MODES[0]:
            key = f"sel_{grade}_{TERMS.index(term)}_{[RANGE_TERM, RANGE_UNTIL, RANGE_ALL].index(rng_how)}"
            st.session_state.setdefault(key, [])
            selected = kanji_picker(key, pool)

    with st.expander("📎 既存プリントを添付（任意・添付した場合はその問題を使います）"):
        up = st.file_uploader("txt / csv / docx / 画像 / PDF",
                              type=["txt", "csv", "md", "docx", "png", "jpg", "jpeg", "pdf"])
        st.caption("txt・csv・docx は「【語|よみ】を含む文」または「語,よみ,例文」の形式で読み取ります。"
                   "画像・PDFは AI が読み取ります（APIキーが必要）。")
        a1, a2 = st.columns(2)
        order = a1.radio("問題の並び順", ORDERS)
        count_mode = a2.radio("問題数", COUNT_MODES)
        rewrite = st.checkbox("答える漢字はそのままで、文を毎回ランダムに作り変える（AI・APIキーが必要）")
        st.caption("文の作り変えでも、答えの漢字・送り仮名・よみは元のプリントのまま変えません。"
                   "学年は添付プリントの学年に合わせてください（文に使う漢字をその学年までに制限します）。")

    ruby = kind in ("ルビ付き", "お助け＋ルビ")

    if st.button("✨ プリントを作成", type="primary"):
        seed = random.randint(0, 10**6)
        st.session_state["seed"] = seed
        items, notes, att = None, [], None
        if up is not None:
            try:
                att = load_attachment(up)
                if not att:
                    notes.append("添付プリントから問題を読み取れませんでした。")
            except Exception as ex:
                notes.append(f"添付ファイルの読み取りに失敗しました: {ex}")
        if att:
            with st.spinner("添付プリントから問題を作成中…"):
                items, nn = build_from_attachment(att, n, count_mode == COUNT_MODES[0],
                                                  order == ORDERS[1], rewrite, grade, seed, ruby)
            notes += nn
        elif pick == PICK_MODES[0] and not selected:
            st.warning("使う漢字を1字以上選んでください。")
        else:
            with st.spinner("問題を作成中…（30秒〜1分ほどかかることがあります）"):
                items, nn = build_items(grade, pool, selected, n, seed, pick != PICK_MODES[0])
            notes += nn
        if items:
            st.session_state["items"] = items
        st.session_state["notes"] = notes

    for msg in st.session_state.get("notes", []):
        st.info(msg)

    items = st.session_state.get("items")
    if not items:
        st.caption("設定と漢字を選んで「プリントを作成」を押してください。")
        return

    scale = SIZES[size_opt] or AUTO_SCALE[grade]
    q_html = sheet_html(items, title, kind, False, publisher, term, scale)
    a_html = sheet_html(items, title, kind, True, publisher, term, scale)
    w, h, _ = PAPERS[paper]
    page_px = int(h * 3.78) + 24

    tab_q, tab_a, tab_b = st.tabs(["問題のみ", "解答のみ", "問題＋解答"])
    with tab_q:
        components.html(document_html([q_html], paper), height=page_px + 70, scrolling=True)
    with tab_a:
        components.html(document_html([a_html], paper), height=page_px + 70, scrolling=True)
    with tab_b:
        components.html(document_html([q_html, a_html], paper), height=page_px * 2 + 70, scrolling=True)

    st.download_button("📥 HTMLを保存（ブラウザで開いて印刷）",
                       document_html([q_html, a_html], paper),
                       file_name="kanji_test.html", mime="text/html")


if __name__ == "__main__":
    main()
